"""
End-to-end tests for the CV Parser & LinkedIn Enrichment System.
Tests the full pipeline: file upload → text extraction → Gemini parsing → DB storage → API response.
External services (Gemini, LinkedIn/Selenium) are mocked.
"""
import io
import json
import os
import tempfile
from unittest.mock import patch, MagicMock

from django.test import TestCase, override_settings
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

from candidates.models import Candidate, Education, Experience, Skill
from candidates.services.file_handler import extract_text
from candidates.services.cv_parser import parse_cv


# Sample parsed CV data that Gemini would return
MOCK_GEMINI_RESPONSE = {
    "full_name": "Yash Lohade",
    "email": "yash@example.com",
    "phone": "+91 9999999999",
    "linkedin_url": "https://linkedin.com/in/yashlohade",
    "summary": "Full-stack developer with 1+ years of experience",
    "location": "Hyderabad, India",
    "current_title": "Software Developer",
    "current_company": "Win Information Technology",
    "skills": [
        {"name": "Python", "proficiency": "advanced"},
        {"name": "Django", "proficiency": "advanced"},
        {"name": "React", "proficiency": "intermediate"},
    ],
    "education": [
        {
            "institution": "XYZ University",
            "degree": "B.Tech",
            "field": "Computer Science",
            "start": "2019",
            "end": "2023",
        }
    ],
    "experience": [
        {
            "company": "Win Information Technology",
            "title": "Software Developer",
            "description": "Backend development with Java and Python",
            "start": "Aug 2025",
            "end": "Present",
        }
    ],
}


def create_test_pdf():
    """Create a minimal valid PDF with resume-like text."""
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    text = """Yash Lohade
yash@example.com | +91 9999999999
LinkedIn: https://linkedin.com/in/yashlohade
Hyderabad, India

Summary:
Full-stack developer with 1+ years of experience in Python, Django, React.

Experience:
Software Developer at Win Information Technology (Aug 2025 - Present)
- Backend development with Java and Python

Education:
B.Tech in Computer Science, XYZ University (2019 - 2023)

Skills: Python, Django, React, PostgreSQL, Docker
"""
    page.insert_text((72, 72), text, fontsize=11)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_test_docx():
    """Create a minimal valid DOCX with resume text."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Yash Lohade")
    doc.add_paragraph("yash@example.com | +91 9999999999")
    doc.add_paragraph("LinkedIn: https://linkedin.com/in/yashlohade")
    doc.add_paragraph("")
    doc.add_paragraph("Summary: Full-stack developer with 1+ years of experience")
    doc.add_paragraph("")
    doc.add_paragraph("Experience:")
    doc.add_paragraph("Software Developer at Win Information Technology (Aug 2025 - Present)")
    doc.add_paragraph("")
    doc.add_paragraph("Education:")
    doc.add_paragraph("B.Tech in Computer Science, XYZ University (2019 - 2023)")
    doc.add_paragraph("")
    doc.add_paragraph("Skills: Python, Django, React")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class FileHandlerTest(TestCase):
    """Test PDF and DOCX text extraction."""

    def test_extract_pdf(self):
        pdf_bytes = create_test_pdf()
        with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as f:
            f.write(pdf_bytes)
            f.flush()
            text = extract_text(f.name)
        os.unlink(f.name)
        self.assertIn('Yash Lohade', text)
        self.assertIn('yash@example.com', text)
        self.assertIn('Python', text)

    def test_extract_docx(self):
        docx_bytes = create_test_docx()
        with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as f:
            f.write(docx_bytes)
            f.flush()
            text = extract_text(f.name)
        os.unlink(f.name)
        self.assertIn('Yash Lohade', text)
        self.assertIn('yash@example.com', text)

    def test_unsupported_format(self):
        with self.assertRaises(ValueError):
            extract_text('/tmp/test.txt')


class CVParserTest(TestCase):
    """Test Gemini CV parsing (mocked)."""

    @patch('candidates.services.cv_parser.genai')
    def test_parse_cv_returns_structured_data(self, mock_genai):
        mock_model = MagicMock()
        mock_genai.GenerativeModel.return_value = mock_model
        mock_response = MagicMock()
        mock_response.text = json.dumps(MOCK_GEMINI_RESPONSE)
        mock_model.generate_content.return_value = mock_response

        result = parse_cv("Some resume text here")

        self.assertEqual(result['full_name'], 'Yash Lohade')
        self.assertEqual(result['email'], 'yash@example.com')
        self.assertEqual(len(result['skills']), 3)
        self.assertEqual(len(result['education']), 1)
        self.assertEqual(len(result['experience']), 1)
        mock_model.generate_content.assert_called_once()

    @patch('candidates.services.cv_parser.genai')
    def test_parse_cv_handles_markdown_fenced_json(self, mock_genai):
        mock_model = MagicMock()
        mock_genai.GenerativeModel.return_value = mock_model
        mock_response = MagicMock()
        mock_response.text = '```json\n' + json.dumps(MOCK_GEMINI_RESPONSE) + '\n```'
        mock_model.generate_content.return_value = mock_response

        result = parse_cv("Resume text")
        self.assertEqual(result['full_name'], 'Yash Lohade')


MEDIA_ROOT_TMP = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=MEDIA_ROOT_TMP)
class UploadAPITest(TestCase):
    """End-to-end test: upload CV → parse → store → return JSON."""

    def setUp(self):
        self.client = APIClient()

    @patch('candidates.views.enrich_candidate')
    @patch('candidates.views.parse_cv')
    def test_upload_pdf_creates_candidate(self, mock_parse, mock_enrich):
        mock_parse.return_value = MOCK_GEMINI_RESPONSE
        mock_enrich.return_value = True

        pdf_bytes = create_test_pdf()
        upload_file = SimpleUploadedFile('resume.pdf', pdf_bytes, content_type='application/pdf')

        response = self.client.post(
            '/api/candidates/upload/',
            {'file': upload_file, 'source_email': 'recruiter@company.com'},
            format='multipart',
        )

        self.assertEqual(response.status_code, 201)
        data = response.json()
        self.assertEqual(data['full_name'], 'Yash Lohade')
        self.assertEqual(data['email'], 'yash@example.com')
        self.assertEqual(data['current_title'], 'Software Developer')
        self.assertEqual(data['source_email'], 'recruiter@company.com')
        self.assertEqual(len(data['skills']), 3)
        self.assertEqual(len(data['education']), 1)
        self.assertEqual(len(data['experience']), 1)

        # Verify DB records
        self.assertEqual(Candidate.objects.count(), 1)
        candidate = Candidate.objects.first()
        self.assertEqual(candidate.full_name, 'Yash Lohade')
        self.assertEqual(Skill.objects.filter(candidate=candidate).count(), 3)
        self.assertEqual(Education.objects.filter(candidate=candidate).count(), 1)
        self.assertEqual(Experience.objects.filter(candidate=candidate).count(), 1)

        # LinkedIn enrichment was attempted
        mock_enrich.assert_called_once_with(candidate)

    @patch('candidates.views.enrich_candidate')
    @patch('candidates.views.parse_cv')
    def test_upload_docx_creates_candidate(self, mock_parse, mock_enrich):
        mock_parse.return_value = MOCK_GEMINI_RESPONSE
        mock_enrich.return_value = True

        docx_bytes = create_test_docx()
        upload_file = SimpleUploadedFile('resume.docx', docx_bytes,
                                         content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        response = self.client.post(
            '/api/candidates/upload/',
            {'file': upload_file},
            format='multipart',
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()['full_name'], 'Yash Lohade')

    def test_upload_unsupported_file_returns_400(self):
        upload_file = SimpleUploadedFile('notes.txt', b'hello', content_type='text/plain')
        response = self.client.post(
            '/api/candidates/upload/',
            {'file': upload_file},
            format='multipart',
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Unsupported file type', response.json()['error'])

    def test_upload_no_file_returns_400(self):
        response = self.client.post('/api/candidates/upload/', {}, format='multipart')
        self.assertEqual(response.status_code, 400)

    @patch('candidates.views.enrich_candidate')
    @patch('candidates.views.parse_cv')
    def test_upload_without_linkedin_skips_enrichment(self, mock_parse, mock_enrich):
        no_linkedin = MOCK_GEMINI_RESPONSE.copy()
        no_linkedin['linkedin_url'] = ''
        mock_parse.return_value = no_linkedin

        pdf_bytes = create_test_pdf()
        upload_file = SimpleUploadedFile('resume.pdf', pdf_bytes, content_type='application/pdf')
        response = self.client.post(
            '/api/candidates/upload/', {'file': upload_file}, format='multipart'
        )

        self.assertEqual(response.status_code, 201)
        mock_enrich.assert_not_called()


@override_settings(MEDIA_ROOT=MEDIA_ROOT_TMP)
class CandidateCRUDTest(TestCase):
    """Test the REST API CRUD endpoints."""

    def setUp(self):
        self.client = APIClient()
        self.candidate = Candidate.objects.create(
            full_name='Test User',
            email='test@example.com',
            phone='+91 1234567890',
            current_title='Developer',
        )
        Skill.objects.create(candidate=self.candidate, name='Python', proficiency_level='advanced')
        Education.objects.create(
            candidate=self.candidate, institution='Test Uni', degree='B.Tech'
        )
        Experience.objects.create(
            candidate=self.candidate, company='Test Corp', title='Dev'
        )

    def test_list_candidates(self):
        response = self.client.get('/api/candidates/')
        self.assertEqual(response.status_code, 200)
        results = response.json()['results']
        self.assertEqual(len(results), 1)
        self.assertEqual(results[0]['full_name'], 'Test User')

    def test_get_candidate_detail(self):
        response = self.client.get(f'/api/candidates/{self.candidate.id}/')
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data['full_name'], 'Test User')
        self.assertEqual(len(data['skills']), 1)
        self.assertEqual(len(data['education']), 1)
        self.assertEqual(len(data['experience']), 1)

    def test_update_candidate(self):
        response = self.client.patch(
            f'/api/candidates/{self.candidate.id}/',
            {'current_title': 'Senior Developer'},
            format='json',
        )
        self.assertEqual(response.status_code, 200)
        self.candidate.refresh_from_db()
        self.assertEqual(self.candidate.current_title, 'Senior Developer')

    def test_delete_candidate(self):
        response = self.client.delete(f'/api/candidates/{self.candidate.id}/')
        self.assertEqual(response.status_code, 204)
        self.assertEqual(Candidate.objects.count(), 0)

    def test_get_nonexistent_returns_404(self):
        response = self.client.get('/api/candidates/9999/')
        self.assertEqual(response.status_code, 404)


class ModelStrTest(TestCase):
    """Test model __str__ methods."""

    def test_candidate_str(self):
        c = Candidate(full_name='John Doe')
        self.assertEqual(str(c), 'John Doe')

    def test_education_str(self):
        c = Candidate.objects.create(full_name='Test')
        e = Education(candidate=c, institution='MIT', degree='MS')
        self.assertEqual(str(e), 'MS - MIT')

    def test_experience_str(self):
        c = Candidate.objects.create(full_name='Test')
        exp = Experience(candidate=c, company='Google', title='SWE')
        self.assertEqual(str(exp), 'SWE at Google')

    def test_skill_str(self):
        c = Candidate.objects.create(full_name='Test')
        s = Skill(candidate=c, name='Python')
        self.assertEqual(str(s), 'Python')
